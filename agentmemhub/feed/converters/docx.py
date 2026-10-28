# -*- coding: utf-8 -*-
"""DOCX 直通：**.docx 不属于"PDF 那一类"**（计划 §7.0.1 专门纠正过这个误判）。

为什么它和 T1 同级，而不是 T3
============================
`.docx` 本质是 ZIP + XML：标题是 `<w:pStyle>`、大纲级别是 `<w:outlineLvl>`、
表格是 `<w:tbl>`、编号是 `<w:numPr>`、图片有 `<a:blip>` 的 rId ——
**语义结构本来就在文件里**，`python-docx` 是确定性抽取：零 LLM、零猜测、可复现。

对比 PDF：PDF 里只有"这个字画在坐标 (x, y)"，标题和表格全靠**重建**，
所以两者在 §7.0 的四个维度上（可及性 / 可提取性 / 保真度 / **可复现性**）都不同级。
把 `.docx` 按 PDF 估工期，是这个计划早期最贵的一个误判。

一个必须避开的坑：段落与表格在 XML 里是**兄弟节点**且交错
====================================================
`document.paragraphs` 与 `document.tables` 是**两个独立列表** —— 分别遍历会让
所有表格挤到文末，文档顺序全丢。必须按 `w:body` 的子节点顺序走
（`iter_block_items`，python-docx 官方 FAQ 的做法）。夹具 `sample.docx` 里
"表格前后的段落"就是这个坑的哨兵：只要顺序错了，那条断言立刻红。

另外两处如实记账（不静默）
======================
· **合并单元格**：markdown 管道表表达不了跨行/跨列，同一格里的话会重复出现 ⇒ 记 warning；
· **图片**：docx 有真实位置（`w:drawing` 就在某个段落里），所以图片**插回原位**，
  比 PDF 那一档强 —— 但正因如此，"抽不出来"才更要报出来。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from ..base import BaseConverter, FeedDocument, FeedError
from ..mdtable import table_md
from ..titles import (TITLE_MAX, clean_meta_title, clip_title, is_junk_title,
                      looks_like_title)
from .office import doc_time_of

#: 认领的扩展名（`.docm` 与 `.docx` 同构：都是 ZIP+XML，只是 content-type 标了宏）
EXTS = (".docx", ".docm")
#: Word 97-2003 的 `.doc` 是 OLE 复合二进制文档，**根本不是 ZIP** ——
#: 硬试只会得到"文件损坏"这种误导性错误，所以显式拦住并给可执行的指引。
LEGACY_EXTS = (".doc",)

#: 内嵌图片上限（一份文档几百张图的极端情况不该把投喂拖成批处理）
MAX_IMAGES = 40

_HEADING_RE = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.I)
_LIST_RE = re.compile(r"^(?:list|列表)", re.I)


def _qn(tag: str):
    from docx.oxml.ns import qn
    return qn(tag)


def iter_block_items(parent):
    """按**文档顺序**遍历段落与表格（`w:body` 的子节点本身就是交错排列的）。

    `parent` 可以是 `Document`（走 `w:body`）或单元格（走 `w:tc`）——
    后者用于将来展开嵌套表格。
    """
    from docx.document import Document as _Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if isinstance(parent, _Document):
        element = parent.element.body
    else:
        element = parent._tc
    for child in element.iterchildren():
        if child.tag == _qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == _qn("w:tbl"):
            yield Table(child, parent)


def _para_text(p) -> str:
    """段落文字。

    不用 `paragraph.text`：它**漏掉 `<w:hyperlink>` 里的 run**（python-docx 的已知局限），
    而带链接的段落恰恰是最容易出现在知识型文档里的（参考资料、引用）。
    按 XML 顺序取所有 `w:t` 才是"在 Word 里看到的那段字"。
    """
    return "".join(node.text or "" for node in p._p.iter(_qn("w:t"))).strip()


def _heading_level(p) -> int:
    """段落的大纲级别（0 = 不是标题）。

    两条判据：样式名（`Heading 1` / 本地化的 `标题 1`）与 `<w:outlineLvl>`。
    后者才是**本质**判据 —— 自定义样式名千奇百怪，但"这是第几级标题"
    在 Word 里就存在大纲级别这一处。
    """
    try:
        m = _HEADING_RE.match((p.style.name or "").strip())
        if m:
            return int(m.group(1))
    except Exception:                        # noqa: BLE001 —— 样式损坏不该毁掉整篇
        pass
    try:
        pPr = p._p.pPr
        if pPr is not None:
            lvl = pPr.find(_qn("w:outlineLvl"))
            if lvl is not None:
                v = int(lvl.get(_qn("w:val")))
                if 0 <= v <= 8:
                    return v + 1
    except Exception:                        # noqa: BLE001
        pass
    return 0


def _is_list(p) -> bool:
    """是否列表项。优先看 `<w:numPr>`（真正的编号属性），样式名只作兜底。"""
    try:
        pPr = p._p.pPr
        if pPr is not None and pPr.numPr is not None:
            return True
    except Exception:                        # noqa: BLE001
        pass
    try:
        return bool(_LIST_RE.match((p.style.name or "").strip()))
    except Exception:                        # noqa: BLE001
        return False


def _para_images(p, *, put_image, state: dict) -> list[str]:
    """段落里的图片 → markdown 引用（**位置就此保住**：docx 里图片属于某个段落）。"""
    blips = [b for b in p._p.iter(_qn("a:blip"))
             if b.get(_qn("r:embed")) or b.get(_qn("r:link"))]
    if not blips:
        return []
    # 先记"文档里到底有几张图"，**再**看能不能抽 —— 否则"没有写入口"时
    # `state["images"]` 恒为 0，告警条件永远不成立，图片就静默消失了
    state["images_found"] += len(blips)
    if not callable(put_image):
        return []
    refs: list[str] = []
    for blip in blips:
        rId = blip.get(_qn("r:embed")) or blip.get(_qn("r:link"))
        try:
            part = state["document"].part.related_parts[rId]
            blob = part.blob
        except Exception as e:               # noqa: BLE001
            state["warnings"].append("内嵌图片（%s）读取失败：%s" % (rId, str(e)[:50]))
            continue
        if not blob:
            continue
        if state["images"] >= MAX_IMAGES:
            if not state.get("image_capped"):
                state["image_capped"] = True
                state["warnings"].append("内嵌图片超过 %d 张，其余未抽取" % MAX_IMAGES)
            break
        state["images"] += 1
        suffix = Path(str(getattr(part, "partname", ""))).suffix or ".png"
        name = "image-%03d%s" % (state["images"], suffix)
        try:
            refs.append(put_image("docx/%s" % name, blob))
        except Exception as e:               # noqa: BLE001
            state["warnings"].append("图片 %s 保存失败：%s" % (name, str(e)[:50]))
    return refs


def _tc_text(tc) -> str:
    """从一个 `w:tc` 取文字。

    不经过 `_Cell` 包装、也不走 `row.cells` —— 因为我们要按**原始 tc 列表**遍历
    （`row.cells` 会把合并区域展开成重复的格子，见 `_table_rows`）。
    """
    parts: list[str] = []
    for para in tc.findall(_qn("w:p")):
        t = "".join(node.text or "" for node in para.iter(_qn("w:t"))).strip()
        if t:
            parts.append(t)
    return " ".join(parts).strip()


def _table_rows(table) -> tuple[list[list[str]], bool]:
    """表格 → `(二维文本, 是否含合并单元格)`。**行列对齐、合并区域不重复。**

    ⚠️ 不能用 `[[c.text for c in row.cells] for row in table.rows]`：
    `row.cells` 对合并区域会把**同一个格子返回多次** ⇒ 内容在 markdown 里被复制好几遍。
    **真实素材命中**：4 份接口测试报告里 3 份有合并单元格，标题行整行合并 ⇒
    同一个标题在表格里出现 3 次，读起来像文档写错了。

    markdown 管道表表达不了"跨行/跨列"，最接近的呈现是**留空**：
      · 横向合并（`w:gridSpan > 1`）⇒ 内容写第一格，其余补空列；
      · 纵向合并的续格（`w:vMerge` 且不是 `restart`）⇒ 留空。
    读者能由此看出"这几格是一体的"；而重复的文字只会让人以为有 3 列相同内容。

    为什么会话"是否含合并"与"渲染"合成一次遍历：两个判据分开写必然漂移 ——
    一边认为"有合并"、另一边却按无合并渲染，质量分会与实际呈现对不上。
    """
    rows: list[list[str]] = []
    merged = False
    try:
        for tr in table._tbl.tr_lst:
            cells: list[str] = []
            for tc in tr.tc_lst:
                span, continued = 1, False
                tcPr = tc.tcPr
                if tcPr is not None:
                    gs = tcPr.find(_qn("w:gridSpan"))
                    if gs is not None:
                        try:
                            span = max(1, int(gs.get(_qn("w:val")) or 1))
                        except Exception:        # noqa: BLE001
                            span = 1
                    vm = tcPr.find(_qn("w:vMerge"))
                    if vm is not None:
                        # 缺 `w:val` 或 `continue` 都表示"这一格是上面那格的延续"
                        continued = (vm.get(_qn("w:val")) or "").lower() != "restart"
                if span > 1 or continued:
                    merged = True
                cells.append("" if continued else _tc_text(tc))
                cells.extend([""] * (span - 1))
            rows.append(cells)
    except Exception as e:                   # noqa: BLE001
        raise FeedError("表格读取失败：%s" % str(e)[:100], kind="convert",
                        retryable=False)
    return rows, merged


def _collect(document, *, put_image) -> tuple[list, dict]:
    """遍历文档，产出**有序**的块列表与统计。"""
    state: dict = {"document": document, "images": 0, "images_found": 0,
                   "warnings": [], "tables": 0, "merged": 0, "headings": 0,
                   "first_para": ""}
    items: list = []
    first_h1 = ""
    for block in iter_block_items(document):
        cls_name = type(block).__name__
        if cls_name == "Table":
            state["tables"] += 1
            rows, merged = _table_rows(block)
            if merged:
                state["merged"] += 1
            items.append(("table", rows))
            continue
        level = _heading_level(block)
        text = _para_text(block)
        images = _para_images(block, put_image=put_image, state=state)
        if level == 1 and text and not first_h1:
            first_h1 = text
        if text and not state["first_para"]:
            # 首个有文字的段落 —— 标题的**最后一档兜底**（见 `_title_of`）
            state["first_para"] = text
        if level:
            state["headings"] += 1
        if not text and not images:
            continue                         # 空段落（Word 里大量存在）不产出块
        items.append(("para", (level, _is_list(block), text, images)))
    state["first_h1"] = first_h1
    return items, state


def _render(items: list, *, head_title: str) -> str:
    blocks: list[str] = []
    if head_title:
        blocks.append("# %s" % head_title)
    for kind, payload in items:
        if kind == "table":
            tmd = table_md(payload)
            if tmd:
                blocks.append(tmd)
            continue
        level, is_list, text, images = payload
        parts: list[str] = []
        if text:
            if level:
                parts.append("%s %s" % ("#" * min(int(level), 6), text))
            elif is_list:
                parts.append("- %s" % text)
            else:
                parts.append(text)
        parts += ["![](%s)" % rel for rel in images]
        if parts:
            blocks.append("\n\n".join(parts))
    return "\n\n".join(b for b in blocks if b).strip() + "\n"


def _title_of(cp, state: dict, ref: str) -> tuple[str, str]:
    """标题：core `title` → 首个一级标题 → **正文首段** → 文件名。

    为什么必须有"正文首段"这一档（真实素材逼出来的）：一份求职文档
    **完全没用标题样式**（38 段全是正文），core title 为空 —— 于是标题落到了文件名上，
    而那是 WPS 的默认名「文字文稿1」。**默认文件名比正文首段更没有信息量**：
    前者对所有这类文档都一样，后者至少是这份文档的内容。

    但首段要**限长**：首段动辄上百字，那不是标题。限长后还要把"这是替代品"
    记进 warning（由调用方做），不能让人以为它就是文档标题。
    """
    t = clean_meta_title(getattr(cp, "title", "") or "")
    if t:
        return t, "core_title"
    if state.get("first_h1"):
        return state["first_h1"], "first_h1"
    para = (state.get("first_para") or "").strip()
    if looks_like_title(para):
        clipped, was_clipped = clip_title(para)
        return clipped, "first_para_clipped" if was_clipped else "first_para"
    stem = Path(str(ref)).stem
    if is_junk_title(stem):
        return "DOCX 文档", "fallback"       # 四个来源全不可用，如实说出来
    return stem, "stem"


class DocxConverter(BaseConverter):
    """`.docx` 直通（认领顺序同样要**排在通用 URL 之前**，理由同 PDF）。"""

    name = "docx"
    version = "1"
    #: 认领的后缀（**类属性**，供格式白名单从注册表收集 —— 见 `feed/formats.py`）
    EXTS = EXTS
    LEGACY_EXTS = LEGACY_EXTS

    def can_handle(self, *, ref: str, kind: str = "",
                   data: bytes | None = None) -> bool:
        if kind == "docx":
            return True
        if kind and kind not in ("url",):
            return False
        suffix = Path(str(ref)).suffix.lower()
        if suffix in EXTS or suffix in LEGACY_EXTS:
            return True
        # 没有扩展名：ZIP 魔数只是必要条件，还得**真的含 word/document.xml** ——
        # 否则任何 zip 都会被它抢走（.xlsx/.pptx 也是 ZIP）
        if data and bytes(data[:4]) == b"PK\x03\x04":
            try:
                import zipfile
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    names = set(z.namelist())
                return "word/document.xml" in names
            except Exception:                # noqa: BLE001
                return False
        return False

    def convert(self, *, ref: str, data: bytes, meta: dict) -> FeedDocument:
        suffix = Path(str(ref)).suffix.lower()
        if suffix in LEGACY_EXTS:
            raise FeedError(
                "这是 Word 97-2003 的老格式 `.doc`（OLE 二进制文档），不是 `.docx` —— "
                "两者内部结构完全不同，`python-docx` 读不了它。"
                "请用 Word 另存为 `.docx` 后重新投喂",
                kind="unsupported", retryable=False)

        from docx import Document

        try:
            document = Document(io.BytesIO(data))
        except Exception as e:               # noqa: BLE001
            raise FeedError("DOCX 结构无法解析（文件可能损坏，或不是真正的 .docx）：%s"
                            % str(e)[:120], kind="convert", retryable=False)

        items, state = _collect(document, put_image=meta.get("put_image"))
        cp = getattr(document, "core_properties", None)
        title, title_src = _title_of(cp, state, ref)
        # 正文里已经有这个标题时不再补 H1：core title 与首个 H1 常常就是同一句话；
        # 而 `first_para*` 那一档更明显 —— 标题**就是**正文首段，补了必然重复
        head = "" if (title_src in ("first_h1", "first_para", "first_para_clipped")
                      or state.get("first_h1") == title) else title
        md = _render(items, head_title=head)

        warnings = list(state["warnings"])
        quality = 1.0
        # 标题是"替代品"时必须让人知道（同 PDF 通路：像是原文的标题比占位标题更危险）
        if title_src == "first_para_clipped":
            warnings.append("文档没有可用标题（无 core title、无标题样式），"
                            "已用正文首段前 %d 字代替，建议人工确认" % TITLE_MAX)
            quality = min(quality, 0.9)
        elif title_src == "fallback":
            warnings.append("无法确定标题（无 core title、无标题样式、首段与文件名都无"
                            "信息量），已用占位标题「%s」" % title)
            quality = min(quality, 0.8)
        if state["merged"]:
            # 合并单元格现在以"空单元格"呈现（内容不再重复），但跨行/跨列的**语义**
            # 在 markdown 里仍然表达不了 ⇒ 轻降 + 如实说明
            warnings.append("有 %d 张表格含合并单元格，markdown 只能以空单元格表示"
                            "（内容不再重复，但跨行/跨列的合并语义有损）" % state["merged"])
            quality = min(quality, 0.9)
        if state["images_found"] and not state["images"]:
            warnings.append("文档含 %d 张图片，但本次没有可用的图片写入口，未保存"
                            % state["images_found"])
            quality = min(quality, 0.7)

        doc_time, time_src = doc_time_of(cp, meta)
        return FeedDocument(
            doc_id=meta["doc_id"], title=title[:200], markdown=md, kind="docx",
            source_ref=str(ref), doc_time=doc_time, converter=self.stamp,
            quality=round(float(quality), 2), warnings=warnings,
            meta={"blocks": len(items), "tables": state["tables"],
                  "headings": state["headings"], "images_extracted": state["images"],
                  "merged_tables": state["merged"], "title_source": title_src,
                  "doc_time_source": time_src, "chars": len(md)},
        )
