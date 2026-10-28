# -*- coding: utf-8 -*-
"""XLSX · 表格工作簿：把**工作表 + 单元格**还原成 markdown 管道表。

为什么它是"常用文档的最后一件"
=============================
DOCX / PPTX / XLSX 同属 OOXML（ZIP + XML，语义结构就在文件里），
所以处理方式同类：**确定性、零 LLM**，难度接近 T1。办公场景里
"笔记 / 汇报 / 台账"三样正好对应 md / pptx / xlsx。

但有一处**必须用库，不能自己解 XML**：

    单元格的日期在文件里是**两步表达** —— `<v>45000</v>` 配上
    `styles.xml` 里 `numFmtId=14` 才是 `2023-03-15`。自己解析极容易把日期
    当整数原样放出去，而那是**看不出来的错**：表里静静多出一列 `45000`，
    不报错、不告警，等你哪天拿它算账才发现。所以用 openpyxl。

能力边界（明说，不假装）
======================
· **公式**按 `data_only=True` 取**缓存值**（Excel 最后一次算出来的结果）。
  代价：文件生成后再没被 Excel 打开过的公式**没有缓存** ⇒ 那一格为空。
  不再加载一次取公式文本 —— 那是"审计公式"的场景，不是知识库；
· **合并单元格**：只读模式（`read_only=True`）拿不到 `merged_cells`，
  所以合并区只有左上角有值、其余为空。这是**有意的取舍**：合并语义
  （"这个值横跨 3 列"）在 markdown 管道表里没有标准表达，而只读模式换来的是
  "几十万行的表也不炸内存"。DOCX 那档做了合并检测 —— 它按段落流式读，
  本来就没有这个内存约束，两处取舍不同是因为约束不同，不是疏漏；
· **图表 / 透视表 / 条件格式 / 批注**都不是单元格文本，不进正文；
· **`.xls`（Excel 97-2003）与加密工作簿**是 OLE 复合二进制（**不是 ZIP**），
  openpyxl 不支持 ⇒ 显式报 `unsupported` 并给出"另存为 .xlsx"的指引
  （与 `.doc` 同一处置）。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from ..base import BaseConverter, FeedDocument, FeedError
from ..mdtable import MAX_COLS, MAX_ROWS, cell_text, table_md
from ..titles import clean_meta_title, clip_title, is_junk_title
from .office import doc_time_of

#: 认领的扩展名（`.xlsm` / `.xltx` / `.xltm` 与 `.xlsx` 同构，只是 content-type 标了宏/模板）
EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")
#: 单份工作簿最多转译的工作表数（防"两千张表"的台账文件把投喂拖死）
MAX_SHEETS = 50
#: OLE 复合文档魔数：`.xls` 老格式与**加密的 xlsx** 共用它（都以它开头）
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
#: 默认工作表名（`Sheet1` / `Sheet 2` / `工作表1`）—— 它们当标题等于零信息
_DEFAULT_SHEET_RE = re.compile(r"^(?:sheet|工作表)\s*\d*$", re.I)


def _is_default_sheet_name(name: str) -> bool:
    n = (name or "").strip()
    return bool(_DEFAULT_SHEET_RE.match(n)) or is_junk_title(n)


def _rows_of(ws, state: dict) -> list[list[str]]:
    """迭代一张工作表 → 参差的文本行（已被 `cell_text` 转成人可读形态）。

    三处细节都是只读模式的**真实脾气**，不是防御性编程：
      · 行长度由 `dimension` 决定，**尾部常是一长串 `None`**（有人把格式刷到第
        16384 列）⇒ 必须裁掉尾部空格子，否则每行都拖着几十个空洞；
      · `dimension` 也可能**虚高**（格式刷出来的"看起来很大"）⇒ 全空行直接跳过；
      · 行数上限按**迭代到的原始行**计（不是"有效行"）—— 因为迭代本身就有成本，
        一张声明了 20 万行的表必须能在 1000 行处停下，而不是先走完再说。
    """
    rows: list[list[str]] = []
    for raw in ws.iter_rows(values_only=True):
        if state["rows"] >= MAX_ROWS:
            state["row_capped"] = True
            break
        state["rows"] += 1
        cells = [cell_text(v) for v in (raw or ())]
        while cells and not cells[-1]:        # 裁尾部空格子
            cells.pop()
        if not cells or not any(cells):
            continue
        if len(cells) > MAX_COLS:
            state["col_capped"] = True
            cells = cells[:MAX_COLS]
        rows.append(cells)
    return rows


def _title_of(cp, ref: str, sheet_names: list[str]) -> tuple[str, str]:
    """标题：core `title` → 文件名 → 首个**有名字的**工作表 → 兜底。

    为什么工作表名排在文件名之后：表名是**表**的名字（"7月台账"），
    文件才是**这份文档**的名字（"2026 年度经营数据"）—— 后者更像标题。
    但前提是文件名不是 `工作簿1` 这种 Office 默认名（那种情况才用表名）。
    """
    t = clean_meta_title(getattr(cp, "title", "") or "")
    if t:
        return t, "core_title"
    stem = Path(str(ref)).stem
    if stem and not is_junk_title(stem):
        clipped, did = clip_title(stem)
        return clipped, "stem_clipped" if did else "stem"
    for name in sheet_names:
        if name and not _is_default_sheet_name(name):
            clipped, did = clip_title(name)
            return clipped, "sheet_clipped" if did else "sheet"
    return "表格工作簿", "fallback"


class XlsxConverter(BaseConverter):
    """`.xlsx` 工作簿（OOXML，与 `.docx` / `.pptx` 同类：确定性、零 LLM）。"""

    name = "xlsx"
    version = "1"
    #: 认领的后缀（**类属性**，供格式白名单从注册表收集 —— 见 `feed/formats.py`）
    EXTS = EXTS

    def can_handle(self, *, ref: str, kind: str = "",
                   data: bytes | None = None) -> bool:
        if kind == "xlsx":
            return True
        if kind and kind not in ("url",):
            return False
        if Path(str(ref)).suffix.lower() in EXTS:
            return True
        # 没有扩展名：ZIP 里含 `xl/workbook.xml` 才算（别抢 .docx / .pptx / .xmind）
        if data and bytes(data[:4]) == b"PK\x03\x04":
            try:
                import zipfile
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    return "xl/workbook.xml" in set(z.namelist())
            except Exception:                # noqa: BLE001
                return False
        return False

    def convert(self, *, ref: str, data: bytes, meta: dict) -> FeedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError:                  # pragma: no cover —— 依赖在 pyproject 里
            raise FeedError("未安装 openpyxl（`uv sync` 可装上）",
                            kind="unsupported", retryable=False)
        # 先在字节层分开"老格式/加密"与"文件坏了"：两者的处置完全不同
        # （前者要用户另存，后者要用户换文件），混成一句报错等于没说
        if bytes(data[:8]) == _OLE_MAGIC:
            raise FeedError(
                "这是 Excel 97-2003 的老格式（.xls）或**加密工作簿** —— 二者都是 "
                "OLE 复合二进制，不是 OOXML。请用 Excel/WPS 另存为 .xlsx 后再投喂"
                "（加密档与 PDF 的 T3d 同一处置）",
                kind="unsupported", retryable=False)
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:               # noqa: BLE001
            raise FeedError("XLSX 结构无法解析（文件可能损坏，或不是真正的工作簿）：%s"
                            % str(e)[:120], kind="convert", retryable=False)
        try:
            return self._build(wb, ref=ref, meta=meta)
        finally:
            try:
                wb.close()
            except Exception:                # noqa: BLE001 —— 关不掉不该盖住真正的结果
                pass

    def _build(self, wb, *, ref: str, meta: dict) -> FeedDocument:
        """有了 `Workbook` 之后的全部逻辑（拆出来是为了 `convert` 里能 `finally: close`）。"""
        state = {"sheets": 0, "rows": 0, "cells": 0, "empty_sheets": 0,
                 "row_capped": False, "col_capped": False, "sheet_capped": False}
        warns: list[str] = []
        blocks: list[tuple[str, list[list[str]]]] = []
        sheet_names: list[str] = []
        total_rows = 0

        for ws in wb.worksheets:
            if len(blocks) >= MAX_SHEETS:
                state["sheet_capped"] = True
                break
            name = str(getattr(ws, "title", "") or "")
            sheet_names.append(name)
            state["rows"] = 0                # 行上限按**每张表**独立计
            state["row_capped"] = False
            state["col_capped"] = False
            rows = _rows_of(ws, state)
            if not rows:
                state["empty_sheets"] += 1
                continue
            if state["row_capped"]:
                warns.append("工作表「%s」超过 %d 行，只转译了前 %d 行"
                             % (name or "-", MAX_ROWS, MAX_ROWS))
            if state["col_capped"]:
                warns.append("工作表「%s」有超过 %d 列的宽表，超出部分未转译"
                             % (name or "-", MAX_COLS))
            blocks.append((name, rows))
            state["sheets"] += 1
            state["cells"] += sum(len(r) for r in rows)
            total_rows += len(rows)

        if state["sheet_capped"]:
            warns.append("工作簿共 %d 张工作表，超过单份处理上限 %d 张，只转译了前 %d 张"
                         % (len(wb.worksheets), MAX_SHEETS, MAX_SHEETS))
        if not blocks:
            raise FeedError("工作簿里没有任何可读的单元格（%d 张工作表全为空）"
                            % len(wb.worksheets), kind="empty", retryable=False)

        cp = getattr(wb, "properties", None)
        title, title_src = _title_of(cp, ref, sheet_names)
        doc_time, time_src = doc_time_of(cp, meta)
        md = self._render(title, blocks, sheets=len(blocks), rows=total_rows)

        # 截断是**丢数据**，必须降分（不只是记 warning）：0.7 让它落进
        # `feed list --low-quality` 的视野，人一眼能看见"这份不完整"
        truncated = bool(state["row_capped"] or state["col_capped"]
                         or state["sheet_capped"])
        return FeedDocument(
            doc_id=meta["doc_id"], title=title[:200], markdown=md, kind="xlsx",
            source_ref=str(ref), doc_time=doc_time, converter=self.stamp,
            quality=0.7 if truncated else 1.0, warnings=warns,
            meta={"sheets": len(blocks), "rows": total_rows,
                  "cells": state["cells"], "empty_sheets": state["empty_sheets"],
                  "truncated": truncated, "title_source": title_src,
                  "doc_time_source": time_src, "chars": len(md)},
        )

    @staticmethod
    def _render(title: str, blocks: list[tuple[str, list[list[str]]]], *,
                sheets: int, rows: int) -> str:
        """渲染。**单表且表名是默认名时不写 `## 表名`** ——

        与 XMIND 的"单画布省略 `## 画布`"同款取舍：`Sheet1` 这种默认名
        写进正文只是噪声；但如果表名是作者起的（"7月台账"），**一定要写** ——
        那是定位信息，丢了就不知道这段表格说的是什么。
        """
        out: list[str] = ["# %s" % title, "",
                          "> 表格工作簿 · 共 %d 张工作表，%d 行" % (sheets, rows)]
        only_one = len(blocks) == 1
        for name, table_rows in blocks:
            out.append("")
            if not (only_one and _is_default_sheet_name(name)):
                out.append("## %s" % (name or "（未命名工作表）"))
                out.append("")
            out.append(table_md(table_rows))
        return "\n".join(out).strip() + "\n"
