# -*- coding: utf-8 -*-
"""生成 T3（PDF）/ DOCX 的**合成夹具** —— 可复现：删掉产物重跑本脚本即可。

为什么这里用合成样本，而 T2 用的是真实抓取样本
============================================
T3 的难点**不是"能不能解析"而是"派给哪条路"**：同一份 `.pdf` 可能是纯文字、
带线条表格、无文本层的扫描件、加密档，四条路由的实现与失败语义完全不同
（`docs/feed-adapter-plan.md` §7.3）。要把四条路由都覆盖到，需要**四份各自
形态确定的 PDF** —— 真实素材里凑齐这四类很难，而且"判对了没有"这件事
恰恰要求样本形态**已知**。所以夹具手工构造最小可解析结构：体积小（都在几 KB）、
每个字节的意图都清楚、不引入第三方样本的版权问题。

真实素材的端到端验收另做（记录在 `docs/modules/feed.md`），两者互补：
合成夹具守**路由判据**，真实素材守**解析质量**。

⚠️ 别把这里的样本当成"真实世界的代表"：真实 PDF 有字体子集、压缩对象流、
CID 编码的 CJK 字形（本脚本刻意回避 —— 手工拼 CJK 字体成本过高）。
所以夹具覆盖"结构判据"，**不覆盖"中文抽取质量"**，后者只能靠真实素材。

重跑：
    .venv/Scripts/python.exe tests/fixtures/feed/make_fixtures_m4.py
"""
from __future__ import annotations

import io
from pathlib import Path

HERE = Path(__file__).resolve().parent


# ---------------------------------------------------------------------------
# 最小 PDF 构造（手工拼结构，不引入 PDF 生成库 —— 生成库会把夹具变成"库的产物"，
# 一旦换库夹具就跟着变，路由测试的稳定性就没了）
# ---------------------------------------------------------------------------

def _esc(text: str) -> str:
    """PDF 字面量字符串转义（`\\` `(` `)` 是结构字符，必须转义）。"""
    return (str(text).replace("\\", r"\\")
            .replace("(", r"\(").replace(")", r"\)"))


def _stream(body: str) -> bytes:
    raw = body.encode("latin-1")
    return b"<< /Length %d >>\nstream\n" % len(raw) + raw + b"\nendstream"


def _build(objects: list[bytes], *, info_id: int = 0) -> bytes:
    """拼出完整 PDF：对象体 + **真实的 xref 偏移表**。

    xref 不能省：省了就得依赖解析器的容错修复（pypdf 会重建、pdfplumber 未必），
    那样"夹具能读"就成了解析器脾气的函数，不是夹具本身的性质。
    """
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref_off = len(out)
    n = len(objects) + 1
    out += b"xref\n0 %d\n" % n
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    trailer = b"<< /Size %d /Root 1 0 R" % n
    if info_id:
        trailer += b" /Info %d 0 R" % info_id
    trailer += b" >>"
    out += b"trailer\n" + trailer + b"\nstartxref\n%d\n%%%%EOF\n" % xref_off
    return bytes(out)


def text_stream(lines: list[str], *, x: float = 72, y: float = 700,
                size: float = 12, leading: float = 18) -> str:
    """一段"多行文字"的内容流（`Td` 逐行下移，等价于行距）。"""
    parts = ["BT", "/F1 %g Tf" % size, "%g %g Td" % (x, y)]
    for i, line in enumerate(lines):
        if i:
            parts.append("0 -%g Td" % leading)
        parts.append("(%s) Tj" % _esc(line))
    parts.append("ET")
    return "\n".join(parts)


def table_stream(rows: list[list[str]], *, x0: float = 72, y_top: float = 700,
                 col_w: float = 130, row_h: float = 26, size: float = 10,
                 caption: str = "") -> str:
    """带**真实线条**的表格页面。

    为什么要画线：`pdfplumber.find_tables()` 默认策略是 `lines`（按矢量线条
    检测单元格边界），而不是按空白猜列。不画线的"看起来像表格"只是排版巧合，
    检测不到 —— 那测的就不是 T3b 路由，而是我们的想象力。
    """
    n_rows, n_cols = len(rows), len(rows[0])
    width = col_w * n_cols
    y_bot = y_top - row_h * n_rows
    out = ["0.8 w"]
    for c in range(n_cols + 1):                       # 竖线
        x = x0 + c * col_w
        out.append("%g %g m %g %g l S" % (x, y_bot, x, y_top))
    for r in range(n_rows + 1):                       # 横线
        y = y_top - r * row_h
        out.append("%g %g m %g %g l S" % (x0, y, x0 + width, y))
    for r, row in enumerate(rows):                    # 单元格文字（基线贴格底）
        for c, cell in enumerate(row):
            tx = x0 + c * col_w + 6
            ty = y_top - r * row_h - row_h + 8
            out.append("BT /F1 %g Tf %g %g Td (%s) Tj ET"
                       % (size, tx, ty, _esc(cell)))
    if caption:
        out.append("BT /F1 13 Tf %g %g Td (%s) Tj ET"
                   % (x0, y_top + 26, _esc(caption)))
    return "\n".join(out)


def make_pdf(pages: list[str], *, title: str = "", created: str = "D:20260101120000Z") -> bytes:
    """把若干内容流拼成一份带 Info 元数据的 PDF。

    对象编号固定为：1 Catalog / 2 Pages / 3 Font / 4 Info / 5.. 页面与内容流。
    """
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",                                          # Pages：等凑齐 Kids 再回填
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ("<< /Title (%s) /CreationDate (%s) >>"
         % (_esc(title), created)).encode("latin-1"),
    ]
    kids: list[int] = []
    for body in pages:
        page_id = len(objs) + 1
        objs.append(("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                     "/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>"
                     % (page_id + 1)).encode("latin-1"))
        objs.append(_stream(body))
        kids.append(page_id)
    objs[1] = ("<< /Type /Pages /Kids [%s] /Count %d >>"
               % (" ".join("%d 0 R" % k for k in kids), len(kids))).encode("latin-1")
    return _build(objs, info_id=4)


# ---------------------------------------------------------------------------
# 四份 PDF 夹具 + 一份 DOCX 夹具
# ---------------------------------------------------------------------------

def build_pdf_text() -> bytes:
    """T3a：两页**纯文字**（每页远超 50 字符，无表格线条）。"""
    page1 = text_stream([
        "Agent Memory Hub - PDF routing fixture",
        "This page exists to exercise the T3a route.",
        "Each page carries well above fifty characters",
        "of extractable text, so the probe must classify",
        "this document as a plain text-layer PDF.",
    ])
    page2 = text_stream([
        "Second page of the same fixture.",
        "Page boundaries must survive conversion,",
        "because the page number is the citation",
        "anchor a human uses to find the passage.",
    ])
    return make_pdf([page1, page2], title="PDF text routing fixture")


def build_pdf_table() -> bytes:
    """T3b：一页 = 表格 + 表题（有线条 ⇒ 表格检测必须命中）。"""
    rows = [
        ["Region", "Revenue", "Growth"],
        ["North", "1280", "+12%"],
        ["South", "940", "-3%"],
        ["West", "1510", "+21%"],
    ]
    body = table_stream(rows, caption="Quarterly numbers by region")
    body += "\n" + text_stream([
        "The paragraph below the table must still be",
        "extracted: a table-aware route that loses the",
        "surrounding prose is a regression, not a win.",
    ], y=520)
    return make_pdf([body], title="PDF table routing fixture")


def _stable_zip(data: bytes, *, date=(2026, 1, 1, 0, 0, 0)) -> bytes:
    """把 ZIP（docx）的条目时间戳固定下来，让产物**逐字节可复现**。

    为什么必须做：`python-docx` 保存时用的是"当前时间"，于是每次生成字节都不同 ——
    夹具在本会话里承诺的"可复现"就成了空话，git 里还会看到没有意义的变更
    （实测：连续生成两次，`sample.docx` 的 SHA256 不同）。
    """
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    try:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                info = zipfile.ZipInfo(item.filename, date_time=date)
                info.compress_type = item.compress_type
                info.external_attr = item.external_attr
                info.internal_attr = item.internal_attr
                info.create_system = item.create_system
                dst.writestr(info, src.read(item.filename))
    finally:
        src.close()
    return out.getvalue()


def build_pdf_scanned() -> bytes:
    """T3c：**图片型** PDF（无文本层）。

    这正是扫描件的本质：**像素里有字，字节里没字**。它同时是 T4 通路的输入形态。

    位图用 Pillow 画，但**手工**把它嵌进最小 PDF 结构，而不是让 Pillow 直接输出 PDF ——
    Pillow 会写进当前时间，夹具字节每次生成都不同（实测确认过），
    那"逐字节可复现"就是一句空话。
    """
    import zlib

    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 320), "white")
    draw = ImageDraw.Draw(img)
    draw.text((30, 40), "SCANNED PAGE - NO TEXT LAYER", fill="black")
    draw.text((30, 90), "Row 1   alpha   beta   gamma", fill="black")
    draw.text((30, 130), "Row 2   delta   epsilon   zeta", fill="black")
    draw.rectangle([20, 20, 880, 300], outline="black", width=2)
    w, h = img.size
    comp = zlib.compress(img.tobytes())

    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",                                            # Pages：稍后回填
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Title () /CreationDate (D:20260101120000Z) >>",
        ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] "
         "/Resources << /XObject << /Im0 7 0 R >> >> /Contents 6 0 R >>"
         % (w, h)).encode("latin-1"),
        _stream("q %d 0 0 %d 0 0 cm /Im0 Do Q" % (w, h)),
        (b"<< /Type /XObject /Subtype /Image /Width %d /Height %d "
         b"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode "
         b"/Length %d >>\nstream\n" % (w, h, len(comp))) + comp + b"\nendstream",
    ]
    objs[1] = b"<< /Type /Pages /Kids [5 0 R] /Count 1 >>"
    return _build(objs, info_id=4)


def build_pdf_encrypted() -> bytes:
    """T3d：**用户密码**加密（`decrypt("")` 必须失败 ⇒ 走 T3d 报错，不硬闯）。"""
    import pypdf

    writer = pypdf.PdfWriter(clone_from=io.BytesIO(build_pdf_text()))
    writer.encrypt("fixture-secret")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_pdf_false_table() -> bytes:
    """T3b 的**误检**守卫：一行文字被线条框成"1 行 2 列"。

    真实素材（求职文档）里 pdfplumber 把一行等宽字体代码检成了 1 行 2 列表格；
    若照单全收，那行字会被当表格**从正文里剔除**，而它又不足以构成表格
    ⇒ 内容两头落空、**永久丢失**（这是静默失真里最坏的一种）。

    构造方式与误检形态一致：三条竖线 + 两条横线围出"一行两列"。
    页面下方另放正常段落，用于验证"没把正文一起削掉"。
    """
    body = table_stream(
        [['const el = driver.findElement(By.id("submit"));',
          'const ok = await el.isVisible();']],
        x0=72, y_top=700, col_w=210, row_h=20)
    body += "\n" + text_stream([
        "The paragraph after the boxed line must survive:",
        "a route that drops prose to gain structure is a regression.",
    ], y=600)
    return make_pdf([body], title="")


def build_pdf_wrapped() -> bytes:
    """首行被**行宽硬切** + **没有元数据标题** —— 标题兜底的守卫。

    真实素材（求职 PDF）里首行是 `…今年有六年有余的工作经`，`验。` 落在第二行；
    直接拿首行当标题就得到一句断掉的话。这里用 ASCII 复现同一形态。
    """
    page = text_stream([
        "Quarterly review of the memory pipeline for the",   # 被行宽切掉
        "north region",                                       # 续行碎片（很短）
        "The remaining lines are ordinary body text and must",
        "not be mistaken for a continuation of the title.",
    ], y=700)
    return make_pdf([page], title="")


def build_docx() -> bytes:
    """DOCX 夹具：**故意混合**标题层级 / 正文 / 列表 / 表格 / 图片。

    混合是有目的的：`.docx` 本质是 ZIP+XML，段落与表格在 XML 里是**兄弟节点**
    （`w:p` 与 `w:tbl` 交错），而 `python-docx` 的 `document.paragraphs` 和
    `document.tables` 是**两个独立列表** —— 分别遍历会把顺序打乱
    （表格全跑到文末）。夹具把这个陷阱摆在同一份文档里，测试才守得住。
    """
    from datetime import datetime

    from docx import Document
    from docx.shared import Inches
    from PIL import Image, ImageDraw

    doc = Document()
    doc.core_properties.title = "投喂适配器 DOCX 夹具"
    doc.core_properties.author = "AgentMemHub"
    # 时间**写死**：夹具必须逐字节可复现（否则每次生成都产生"新内容"，
    # 幂等测试与 doc_time 断言都会随机失败）
    doc.core_properties.modified = datetime(2026, 1, 1, 12, 0, 0)

    doc.add_heading("投喂适配器 DOCX 夹具", level=1)
    doc.add_paragraph("这一段是正文段落，用来验证段落被原样保留。")
    doc.add_heading("一、表格前后各有一段", level=2)
    doc.add_paragraph("表格之前的段落。")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    data = [["列一", "列二", "列三"],
            ["甲", "乙", "丙"],
            ["1", "2", "3"]]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.add_paragraph("表格之后的段落 —— 它必须出现在表格**后面**，不能被挤到文末。")
    doc.add_heading("二、列表", level=2)
    doc.add_paragraph("第一项", style="List Bullet")
    doc.add_paragraph("第二项", style="List Bullet")
    doc.add_heading("三、图片", level=2)
    img = Image.new("RGB", (320, 120), "white")
    ImageDraw.Draw(img).text((16, 50), "DOCX IMAGE FIXTURE", fill="black")
    ibuf = io.BytesIO()
    img.save(ibuf, format="PNG")
    ibuf.seek(0)
    doc.add_picture(ibuf, width=Inches(2.4))

    out = io.BytesIO()
    doc.save(out)
    # 过一遍固定时间戳的 ZIP —— 否则每次生成字节都不同（见 `_stable_zip`）
    return _stable_zip(out.getvalue())


#: 产物清单：文件名 → 构造函数
FIXTURES = {
    "pdf_text.pdf": build_pdf_text,
    "pdf_table.pdf": build_pdf_table,
    "pdf_scanned.pdf": build_pdf_scanned,
    "pdf_encrypted.pdf": build_pdf_encrypted,
    "pdf_false_table.pdf": build_pdf_false_table,
    "pdf_wrapped.pdf": build_pdf_wrapped,
    "sample.docx": build_docx,
}


def main() -> int:
    for name, builder in FIXTURES.items():
        data = builder()
        (HERE / name).write_bytes(data)
        print("wrote %-20s %7d bytes" % (name, len(data)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
