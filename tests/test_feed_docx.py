# -*- coding: utf-8 -*-
"""M4 · DOCX 直通：**顺序**是这一档最容易错的地方，其余都是确定性抽取。

夹具 `sample.docx`（`make_fixtures_m4.py` 生成）里刻意把"段落—表格—段落"交错摆放：
`python-docx` 的 `document.paragraphs` 与 `document.tables` 是**两个独立列表**，
分别遍历会让表格全部挤到文末。这条断言就是那个坑的哨兵。

**隔离**：注入 tmp 配置 —— 绝不碰真实的 `database/materials.db` 与 `exports/`。
"""
from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest

from agentmemhub.feed import get_converter
from agentmemhub.feed.converters.docx import DocxConverter
from agentmemhub.feed.service import ingest

FIX = Path(__file__).resolve().parent / "fixtures" / "feed"


def _fx(name: str) -> bytes:
    return (FIX / name).read_bytes()


class _FakeCfg:
    def __init__(self, root: Path):
        self.root = root
        self.paths = {"db": root / "materials.db",
                      "store": root / "store",
                      "kb": root / "kb"}

    def feed_paths(self):
        return dict(self.paths)

    def feed(self):
        return {"convert": {"llm": False, "ocr": "local"}}


@pytest.fixture()
def cfg(tmp_path):
    return _FakeCfg(tmp_path)


def _docx(tmp_path, name="样本.docx", data=None) -> Path:
    p = tmp_path / name
    p.write_bytes(_fx("sample.docx") if data is None else data)
    return p


def _quiet(*_a, **_kw):
    return None


def _convert(data=None, ref="样本.docx", **meta_extra):
    """直接调转换器（不过 service），便于对 meta / 回调做替身注入。"""
    meta = {"doc_id": "testdoc", "mtime": 0, "now": 0}
    meta.update(meta_extra)
    return DocxConverter().convert(ref=ref, data=_fx("sample.docx") if data is None else data,
                                   meta=meta)


# ---------------------------------------------------------------------------
# 顺序：这一档的核心风险
# ---------------------------------------------------------------------------

def test_顺序_表格夹在两段之间不被挤到文末():
    doc = _convert()
    md = doc.markdown
    i_before = md.index("表格之前的段落")
    i_table = md.index("| 列一 | 列二 | 列三 |")
    i_after = md.index("表格之后的段落")
    assert i_before < i_table < i_after, (
        "文档顺序丢了 —— 说明用了 document.paragraphs / document.tables 分别遍历")


def test_顺序_标题层级按文档顺序出现():
    md = _convert().markdown
    assert md.index("一、表格前后各有一段") < md.index("二、列表") < md.index("三、图片")


# ---------------------------------------------------------------------------
# 结构抽取
# ---------------------------------------------------------------------------

def test_标题层级映射为markdown井号():
    md = _convert().markdown
    assert "# 投喂适配器 DOCX 夹具" in md          # Heading 1
    assert "## 一、表格前后各有一段" in md          # Heading 2


def test_列表项渲染为无序列表():
    md = _convert().markdown
    assert "- 第一项" in md and "- 第二项" in md


def test_表格渲染为管道表且行列齐全():
    md = _convert().markdown
    assert "| 列一 | 列二 | 列三 |" in md
    assert "| 甲 | 乙 | 丙 |" in md
    assert "| 1 | 2 | 3 |" in md


def test_标题优先用core_properties的title():
    doc = _convert()
    assert doc.title == "投喂适配器 DOCX 夹具"
    assert doc.meta["title_source"] == "core_title"


def test_标题不许重复出现两次():
    """core title 与首个 H1 是同一句话时，只留正文里那一个 H1（不额外补标题）。"""
    md = _convert().markdown
    assert md.count("投喂适配器 DOCX 夹具") == 1


def test_doc_time取core_modified而不是模板created():
    """python-docx 新建文档的 `created` 是**模板时间**（实测 2013-12-23），
    拿它当"知识产生时间"会让新资料在时间衰减里排到十年前。"""
    doc = _convert()
    assert doc.meta["doc_time_source"] == "core_modified"
    assert doc.doc_time == int(datetime(2026, 1, 1, 12, 0, tzinfo=timezone.utc).timestamp())


# ---------------------------------------------------------------------------
# 图片：docx 有真实位置，所以图要插回原位
# ---------------------------------------------------------------------------

def test_图片落mirror并插回原位():
    seen: dict[str, bytes] = {}

    def _put(url, blob):
        seen[url] = blob
        return "assets/" + url.split("/")[-1]

    doc = _convert(put_image=_put)
    assert doc.meta["images_extracted"] == 1
    assert len(seen) == 1
    url, blob = next(iter(seen.items()))
    assert blob[:4] == b"\x89PNG"
    assert "](assets/image-001.png)" in doc.markdown


def test_图片位置在所属段落之后():
    md = _convert(put_image=lambda url, blob: "assets/" + url.split("/")[-1]).markdown
    # 图片在"三、图片"这个小标题所在的段落里 ⇒ 紧跟其后
    assert md.index("三、图片") < md.index("![](")


def test_没有图片写入口时如实告警不静默():
    doc = _convert()
    assert doc.meta["images_extracted"] == 0
    assert any("图片" in w for w in doc.warnings)
    assert doc.quality < 1.0


# ---------------------------------------------------------------------------
# 质量与端到端
# ---------------------------------------------------------------------------

def test_质量_无合并单元格时满分(tmp_path, cfg):
    r = ingest(str(_docx(tmp_path)), config=cfg, log=_quiet)
    assert r["ok"] is True, r.get("error")
    assert r["kind"] == "docx" and r["converter"] == "docx@1"
    assert r["quality"] == 1.0
    assert r["warnings"] == []


def test_端到端_出kb页且frontmatter带溯源(tmp_path, cfg):
    r = ingest(str(_docx(tmp_path)), config=cfg, log=_quiet)
    text = Path(r["kb_path"]).read_text(encoding="utf-8")
    assert 'origin: external' in text
    assert 'kind: "docx"' in text
    assert "| 甲 | 乙 | 丙 |" in text


def test_合并单元格被如实打标():
    """markdown 管道表表达不了合并单元格 ⇒ 内容会在相邻格重复，必须打标让人复核。"""
    from docx import Document

    d = Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "跨列标题"
    t.cell(1, 0).text = "a"
    t.cell(1, 1).text = "b"
    buf = io.BytesIO()
    d.save(buf)
    doc = DocxConverter().convert(ref="m.docx", data=buf.getvalue(),
                                  meta={"doc_id": "m", "mtime": 0, "now": 0})
    assert any("合并单元格" in w for w in doc.warnings)
    assert doc.quality < 1.0


# ---------------------------------------------------------------------------
# 样式形态：真实 Word 文档的样式名**不确定**，判据必须抓本质
# ---------------------------------------------------------------------------

def _convert_doc(document) -> str:
    buf = io.BytesIO()
    document.save(buf)
    return DocxConverter().convert(ref="c.docx", data=buf.getvalue(),
                                   meta={"doc_id": "c", "mtime": 0, "now": 0}).markdown


def test_中文样式名的标题也能识别():
    """中文版 Word 文档里样式名可能是 `标题 1` 而不是 `Heading 1`。"""
    from docx import Document

    d = Document()
    h1 = d.styles["Heading 1"]
    h1.name = "标题 1"
    d.add_paragraph("中文样式标题", style=h1)
    assert "# 中文样式标题" in _convert_doc(d)


def test_自定义样式靠大纲级别识别为标题():
    """样式名可以千奇百怪，但"这是第几级标题"在 Word 里只存在大纲级别一处。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    p = d.add_paragraph("自定义样式但其实是一级标题")
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "0")
    p._p.get_or_add_pPr().append(lvl)
    assert "# 自定义样式但其实是一级标题" in _convert_doc(d)


def test_超链接里的文字不丢():
    """`paragraph.text` **漏掉 `<w:hyperlink>` 里的 run**（python-docx 已知局限），
    而带链接的段落恰恰最容易出现在知识型文档里（参考资料、引用）。"""
    from docx import Document
    from docx.oxml import OxmlElement

    d = Document()
    p = d.add_paragraph("参考：")
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "项目主页"
    run.append(text)
    link.append(run)
    p._p.append(link)
    assert "项目主页" in _convert_doc(d)


# ---------------------------------------------------------------------------
# 老格式与坏文件：**显式报错并指路**，不是"文件损坏"
# ---------------------------------------------------------------------------

def test_老格式doc给出另存为指引(tmp_path, cfg):
    p = tmp_path / "老文档.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)   # OLE 魔数
    r = ingest(str(p), config=cfg, log=_quiet)
    assert r["ok"] is False
    assert r["error_kind"] == "unsupported"
    assert "另存为" in r["error"] and ".docx" in r["error"]


def test_损坏的docx报convert错误():
    from agentmemhub.feed.base import FeedError

    with pytest.raises(FeedError) as ei:
        _convert(data=b"PK\x03\x04 not a real docx")
    assert ei.value.kind == "convert"


# ---------------------------------------------------------------------------
# 认领
# ---------------------------------------------------------------------------

def test_认领_无扩展名但真是docx():
    conv = get_converter(ref="download", kind="", data=_fx("sample.docx"))
    assert conv is not None and conv.name == "docx"


def test_认领_别的zip不被当成docx():
    """ZIP 魔数只是必要条件：`.xlsx`/`.pptx` 同样是 ZIP，不能都抢过来。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<worksheet/>")
    assert get_converter(ref="download", kind="", data=buf.getvalue()) is None


def test_认领_https链接的docx归docx而非url():
    conv = get_converter(ref="https://example.com/a.docx", kind="url",
                         data=_fx("sample.docx"))
    assert conv is not None and conv.name == "docx"
