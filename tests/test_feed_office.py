# -*- coding: utf-8 -*-
"""M5 · XMIND 与 PPTX：两个"重要文档"格式的适配。

**夹具是现造的**（`_make_xmind` / `_make_pptx`）—— 它们守的是**结构判据**
（层级递归、画布分节、备注/表格/图片的取法），不是"某份真实文件长什么样"。
真实素材的验收记录见 `docs/modules/feed.md`（14 份 xmind / 4 份 pptx 实测）。

**隔离**：全部用 BytesIO 直接喂转换器，不碰真实库与产物目录。
"""
from __future__ import annotations

import io
import json
import zipfile

import pytest

from agentmemhub.feed import get_converter
from agentmemhub.feed.base import FeedError
from agentmemhub.feed.converters.pptx import PptxConverter
from agentmemhub.feed.converters.xmind import XmindConverter


def _quiet(*_a, **_kw):
    return None


def _conv_xmind(data: bytes, ref="导图.xmind", **meta_extra):
    meta = {"doc_id": "x", "mtime": 0, "now": 0}
    meta.update(meta_extra)
    return XmindConverter().convert(ref=ref, data=data, meta=meta)


def _conv_pptx(data: bytes, ref="演示.pptx", **meta_extra):
    meta = {"doc_id": "x", "mtime": 0, "now": 0}
    meta.update(meta_extra)
    return PptxConverter().convert(ref=ref, data=data, meta=meta)


# ---------------------------------------------------------------------------
# XMIND
# ---------------------------------------------------------------------------

def _topic(title, children=None, **extra):
    t = {"title": title}
    if children:
        t["children"] = {"attached": children}
    t.update(extra)
    return t


def _make_xmind(sheets, *, with_json=True, with_xml=True) -> bytes:
    """造一份 xmind（ZIP）。`sheets` 是 `[(画布名, 中心主题 dict)]`。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        if with_json:
            payload = []
            for i, (name, root) in enumerate(sheets):
                item = {"id": "s%d" % i, "class": "sheet", "rootTopic": root}
                if name is not None:
                    item["title"] = name
                payload.append(item)
            z.writestr("content.json", json.dumps(payload, ensure_ascii=False))
        if with_xml:
            z.writestr("content.xml", "<xmap-content/>")
    return buf.getvalue()


def test_xmind_层级还原成嵌套列表():
    data = _make_xmind([(None, _topic("产品规划", [
        _topic("一期", [_topic("登录"), _topic("支付")]),
        _topic("二期"),
    ]))])
    md = _conv_xmind(data).markdown
    assert md.count("- 产品规划") == 1
    assert "  - 一期" in md
    assert "    - 登录" in md and "    - 支付" in md
    assert "  - 二期" in md


def test_xmind_单画布不加画布小节():
    """单画布时导图名就是文档标题，再加一层 `## 画布` 纯属重复
    （与 PDF"单页不加页码"同一条道理）。"""
    md = _conv_xmind(_make_xmind([(None, _topic("只有一张图"))])).markdown
    assert "## " not in md


def test_xmind_多画布按画布分节():
    data = _make_xmind([("画布甲", _topic("甲根")), ("画布乙", _topic("乙根"))])
    doc = _conv_xmind(data)
    assert doc.meta["sheets"] == 2
    md = doc.markdown
    assert "## 画布甲" in md and "## 画布乙" in md
    assert any("画布" in w for w in doc.warnings), "多画布要说明跨画布关系未保留"


def test_xmind_默认画布名回落到中心主题():
    """XMind 默认画布名是「思维导图」—— 对所有导图都一样，等于没名字。"""
    md = _conv_xmind(_make_xmind([("思维导图", _topic("真正的名字")),
                                  ("思维导图", _topic("第二张")),])).markdown
    assert "## 真正的名字" in md
    assert "## 思维导图" not in md


def test_xmind_默认根节点名不当标题():
    """「中心主题」是 XMind 新建导图的默认根名（实测 2 份业务导图没改过）⇒ 回落文件名。"""
    doc = _conv_xmind(_make_xmind([(None, _topic("中心主题", [_topic("子项")]))]),
                      ref="APK整改.xmind")
    assert doc.title == "APK整改"
    assert doc.meta["title_source"] == "stem"


def test_xmind_备注与标签压成行内():
    """备注在思维导图里常是"为什么这么决策" —— 丢了就只剩结论。"""
    data = _make_xmind([(None, _topic("决策", [
        _topic("方案 A", notes={"plain": {"content": "因为时延低"}}, labels=["已定"]),
    ]))])
    md = _conv_xmind(data).markdown
    assert "因为时延低" in md
    assert "`已定`" in md
    # 必须**压在同一行内** —— 换行会打断嵌套列表的层级
    line = [ln for ln in md.splitlines() if "方案 A" in ln][0]
    assert "因为时延低" in line and "`已定`" in line


def test_xmind_优先读content_json而不是旧版空壳xml():
    """**实测教训**：用户 14 份导图里 `content.xml` 全是同一份 4,309 字节的空壳
    （旧版兼容留的），先读它只会得到一份空导图。"""
    data = _make_xmind([(None, _topic("真数据", [_topic("子")]))])
    doc = _conv_xmind(data)
    assert doc.meta["source"] == "content.json"
    assert "真数据" in doc.markdown


def test_xmind_只有旧版xml时也能读():
    xml = ('<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0">'
           '<sheet id="s1"><title>画布一</title>'
           '<topic id="t1"><title>老版根</title><children>'
           '<topics type="attached"><topic id="t2"><title>老版子项</title></topic>'
           '</topics></children></topic></sheet></xmap-content>').encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("content.xml", xml)
    doc = _conv_xmind(buf.getvalue())
    assert doc.meta["source"] == "content.xml"
    assert "老版根" in doc.markdown and "老版子项" in doc.markdown


def test_xmind_空导图如实报错():
    data = _make_xmind([(None, _topic(""))])
    with pytest.raises(FeedError) as ei:
        _conv_xmind(data)
    assert ei.value.kind == "empty" and ei.value.retryable is False


def test_xmind_认领():
    data = _make_xmind([(None, _topic("x"))])
    assert get_converter(ref="a.xmind", kind="").name == "xmind"
    assert get_converter(ref="download", kind="", data=data).name == "xmind"
    # 别的 ZIP（docx）不能被 xmind 抢走
    assert get_converter(ref="a.docx", kind="").name == "docx"


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _make_pptx(*, slides=1, title=None, with_notes=False, with_table=False,
               with_picture=False, set_core_title=None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]                      # 空白版式
    for i in range(slides):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = "第 %d 页的正文" % (i + 1)
        if with_table and i == 0:
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3),
                                           Inches(4), Inches(1))
            table = shape.table
            table.cell(0, 0).text = "列一"
            table.cell(0, 1).text = "列二"
            table.cell(1, 0).text = "甲"
            table.cell(1, 1).text = "乙"
        if with_picture and i == 0:
            from PIL import Image
            ib = io.BytesIO()
            Image.new("RGB", (60, 40), "white").save(ib, format="PNG")
            ib.seek(0)
            slide.shapes.add_picture(ib, Inches(1), Inches(5), Inches(1))
        if with_notes and i == 0:
            slide.notes_slide.notes_text_frame.text = "这是讲稿备注"
    if set_core_title:
        prs.core_properties.title = set_core_title
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_pptx_按页分节且页码是标题的一部分():
    doc = _conv_pptx(_make_pptx(slides=3))
    md = doc.markdown
    assert "## 第 1 页" in md and "## 第 2 页" in md and "## 第 3 页" in md
    assert "第 2 页的正文" in md
    assert doc.meta["slides"] == 3 and doc.meta["text_pages"] == 3


def test_pptx_备注一定要取到():
    """**实测**：评审 PPT 的备注里写的是讲稿与补充说明，比正文更有信息量。"""
    md = _conv_pptx(_make_pptx(with_notes=True)).markdown
    assert "这是讲稿备注" in md
    assert "> 备注：" in md


def test_pptx_表格走统一表格渲染():
    doc = _conv_pptx(_make_pptx(with_table=True))
    assert doc.meta["tables"] == 1
    assert "| 列一 | 列二 |" in doc.markdown
    assert "| 甲 | 乙 |" in doc.markdown


def test_pptx_图片落mirror且不被当成位置信息():
    seen: dict[str, bytes] = {}

    def put_image(url, blob):
        seen[url] = blob
        return "assets/" + url.split("/")[-1]

    doc = _conv_pptx(_make_pptx(with_picture=True), put_image=put_image)
    assert doc.meta["images_extracted"] == 1
    assert seen and next(iter(seen.values()))[:4] == b"\x89PNG"
    assert "![](" in doc.markdown
    assert not doc.warnings, "普通形状不该被误报成「图片读取失败」"


def test_pptx_普通形状不误报为图片失败():
    """**实测踩到**：早先对每个"没有文字的形状"都试 `shape.image` ——
    线条/自选图形/空文本框会抛异常，于是 8 页的 PPT 报出"8 张图片读取失败"
    （实际只有 4 张图）。**假告警比漏报更伤。**"""
    doc = _conv_pptx(_make_pptx(slides=2, with_picture=False),
                     put_image=lambda u, b: "assets/x.png")
    assert doc.meta["images_extracted"] == 0
    assert not any("图片读取失败" in w for w in doc.warnings), doc.warnings


def test_pptx_默认core标题不当标题():
    """「PowerPoint 演示文稿」是 PowerPoint 的默认 core title（实测 4 份 PPT 全是它）。"""
    doc = _conv_pptx(_make_pptx(set_core_title="PowerPoint 演示文稿"),
                     ref="软工职级评审-谢睦邻.pptx")
    assert doc.title == "软工职级评审-谢睦邻"
    assert doc.meta["title_source"] == "stem"


def test_pptx_认领与不抢别人():
    data = _make_pptx()
    assert get_converter(ref="a.pptx", kind="").name == "pptx"
    assert get_converter(ref="download", kind="", data=data).name == "pptx"
    # docx / xmind 都是 ZIP，不能被 pptx 抢走
    xm = _make_xmind([(None, _topic("x"))])
    assert get_converter(ref="download", kind="", data=xm).name == "xmind"


def test_pptx_损坏文件报convert错误():
    with pytest.raises(FeedError) as ei:
        _conv_pptx(b"PK\x03\x04 not really a pptx")
    assert ei.value.kind == "convert" and ei.value.retryable is False
